from pathlib import Path
from datetime import date
from PIL import Image, ImageDraw, ImageFont

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path("/Users/evgeniyloshak/Documents/ECOM")
WORK = ROOT / ".avdar_report_work"
WORK.mkdir(parents=True, exist_ok=True)
OUT = ROOT / "Avdar_Market_Deep_Research_2026-09-14.docx"

NAVY = "203A5F"
BLUE = "2E5B88"
GREEN = "3E6B5B"
AMBER = "C58B2A"
RED = "A64B4B"
LIGHT_BLUE = "EDF3F8"
PALE = "F6F7F9"
MID = "6B7280"
GRID = "D9D9D9"
BLACK = "000000"


def set_run_font(run, name="Arial", size=None, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=GRID, size="5"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = tcBorders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tcBorders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_row_cant_split(row):
    """Keep an entire table row on one page when Word paginates the document."""
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))


def set_table_no_autofit(table):
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_paragraph_keep(paragraph, with_next=False, together=False):
    pPr = paragraph._p.get_or_add_pPr()
    if with_next:
        pPr.append(OxmlElement("w:keepNext"))
    if together:
        pPr.append(OxmlElement("w:keepLines"))


def remove_paragraph_borders(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    for border in list(pPr.findall(qn("w:pBdr"))):
        pPr.remove(border)


def add_hyperlink(paragraph, text, url, color=BLUE, underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    rPr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Arial")
    rFonts.set(qn("w:hAnsi"), "Arial")
    rPr.append(rFonts)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = " PAGE "
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    set_run_font(run, size=8, color=MID)


def add_para(doc, text="", *, bold_lead=None, style=None, before=0, after=5, line=1.12, keep=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    if keep:
        set_paragraph_keep(p, together=True)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.left_indent = Inches(0.22 + 0.20 * level)
        p.paragraph_format.first_line_indent = Inches(-0.14)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.08
        if isinstance(item, tuple):
            lead, rest = item
            r = p.add_run(lead)
            set_run_font(r, bold=True)
            r = p.add_run(rest)
            set_run_font(r)
        else:
            r = p.add_run(item)
            set_run_font(r)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.24)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(item)
        set_run_font(r)


def add_table(doc, headers, rows, widths=None, font_size=8.8, header_fill=NAVY, first_col_left=True, caption=None):
    if caption:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(caption)
        set_run_font(r, size=9, bold=True, color=BLACK)
        set_paragraph_keep(p, with_next=True)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_no_autofit(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_row_cant_split(hdr)
    for j, text in enumerate(headers):
        cell = hdr.cells[j]
        if widths:
            cell.width = Inches(widths[j])
        set_cell_shading(cell, header_fill)
        set_cell_borders(cell)
        set_cell_margins(cell, 90, 90, 90, 90)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(str(text))
        set_run_font(r, size=font_size, bold=True, color="FFFFFF")
    for i, row in enumerate(rows):
        body_row = table.add_row()
        set_row_cant_split(body_row)
        cells = body_row.cells
        fill = "FFFFFF" if i % 2 == 0 else PALE
        for j, text in enumerate(row):
            cell = cells[j]
            if widths:
                cell.width = Inches(widths[j])
            set_cell_shading(cell, fill)
            set_cell_borders(cell)
            set_cell_margins(cell, 80, 90, 80, 90)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if j == 0 and first_col_left:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(text)) < 34 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(text))
            set_run_font(r, size=font_size, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(13 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5 if level == 1 else 4)
    for r in p.runs:
        set_run_font(r, size=18 if level == 1 else 13.5, bold=True, color=BLACK)
    return p


def page_break(doc):
    doc.add_page_break()


def make_charts():
    font_dir = Path("/System/Library/Fonts/Supplemental")
    regular_path = font_dir / "Arial.ttf"
    bold_path = font_dir / "Arial Bold.ttf"

    def font(size, bold=False):
        return ImageFont.truetype(str(bold_path if bold else regular_path), size)

    def text_center(draw, box, label, text_font, fill):
        left, top, right, bottom = box
        bounds = draw.textbbox((0, 0), label, font=text_font)
        width = bounds[2] - bounds[0]
        height = bounds[3] - bounds[1]
        draw.text(
            ((left + right - width) / 2, (top + bottom - height) / 2 - bounds[1]),
            label,
            font=text_font,
            fill=fill,
        )

    def save_chart(image, path):
        image.save(path, format="PNG", dpi=(220, 220), optimize=True)

    # Market score bars, 7.1 × 3.6 inches at 220 dpi.
    markets = ["Германия", "Нидерланды", "Бельгия", "Дания", "Швеция", "Польша", "Финляндия"]
    scores = [90, 89, 81, 78, 76, 73, 55]
    colors = ["#203A5F", "#2E5B88", "#3E6B5B", "#6F8FAF", "#8EA5B9", "#C58B2A", "#A64B4B"]
    image = Image.new("RGB", (1562, 792), "white")
    draw = ImageDraw.Draw(image)
    draw.text((20, 18), "Рейтинг рынков остаётся без изменений", font=font(40, True), fill="#000000")
    left, right, top, bottom = 330, 1460, 105, 700
    for tick in range(0, 101, 20):
        x = left + (right - left) * tick / 100
        draw.line((x, top, x, bottom), fill="#E5E7EB", width=2)
        label = str(tick)
        bounds = draw.textbbox((0, 0), label, font=font(23))
        draw.text((x - (bounds[2] - bounds[0]) / 2, bottom + 10), label, font=font(23), fill="#6B7280")
    step = (bottom - top) / len(markets)
    for i, (market, score, color) in enumerate(zip(markets, scores, colors)):
        cy = top + step * (i + 0.5)
        draw.text((20, cy - 17), market, font=font(28), fill="#111827")
        bar_end = left + (right - left) * score / 100
        draw.rounded_rectangle((left, cy - 24, bar_end, cy + 24), radius=8, fill=color)
        draw.text((bar_end + 14, cy - 16), str(score), font=font(27, True), fill="#111827")
    text_center(draw, (left, bottom + 38, right, 790), "Баллы из 100", font(24), "#4B5563")
    save_chart(image, WORK / "market_scores.png")

    # Inventory quantity and landed-value bars, 7.15 × 4.2 inches at 220 dpi.
    labels = ["Пледы овечья шерсть", "Пледы кашемир", "Длинные носки", "Шарфы кашемир", "Плед из яка", "Короткие носки", "Camel шаль", "Шарфы шерсть", "Шали шерсть", "Перчатки"]
    qty = [33, 6, 121, 49, 5, 113, 7, 28, 5, 13]
    value = [2290.87, 1012.30, 838.53, 787.05, 515.15, 508.50, 233.58, 227.70, 178.39, 128.70]
    image = Image.new("RGB", (1573, 924), "white")
    draw = ImageDraw.Draw(image)
    draw.text((18, 16), "Штуки сосредоточены в носках, капитал — в пледах", font=font(38, True), fill="#000000")
    label_left, chart1_left, chart1_right = 20, 405, 900
    chart2_left, chart2_right = 1030, 1545
    top, bottom = 155, 875
    draw.text((chart1_left, 90), "Количество единиц", font=font(30, True), fill="#111827")
    draw.text((chart2_left, 90), "Landed value в евро", font=font(30, True), fill="#111827")
    for x0, x1, maximum in ((chart1_left, chart1_right, 130), (chart2_left, chart2_right, 2500)):
        for tick in range(0, 6):
            x = x0 + (x1 - x0) * tick / 5
            draw.line((x, top, x, bottom), fill="#E5E7EB", width=2)
    step = (bottom - top) / len(labels)
    for i, (label, count, amount) in enumerate(zip(labels, qty, value)):
        cy = top + step * (i + 0.5)
        draw.text((label_left, cy - 14), label, font=font(23), fill="#111827")
        q_end = chart1_left + (chart1_right - chart1_left) * count / 130
        v_end = chart2_left + (chart2_right - chart2_left) * amount / 2500
        draw.rectangle((chart1_left, cy - 18, q_end, cy + 18), fill="#2E5B88")
        draw.rectangle((chart2_left, cy - 18, v_end, cy + 18), fill="#C58B2A")
    save_chart(image, WORK / "inventory.png")

    # Google Trends matrix, 7.0 × 4.8 inches at 220 dpi.
    families = ["Cashmere scarf", "Wool scarf", "Wool socks", "Yak socks", "Wool gloves", "Camel shawl", "Wool shawl", "Wool throw", "Yak throw", "Cashmere throw", "Gift exact"]
    values = [
        [2, 1, 0, 2, 1, 0, 2],
        [3, 2, 1, 1, 1, 1, 2],
        [3, 3, 2, 2, 3, 3, 2],
        [0, 0, 0, 0, 0, 0, 0],
        [2, 1, 0, 0, 0, 0, 2],
        [0, 0, 0, 0, 0, 0, 0],
        [1, 0, 0, 0, 0, 0, 1],
        [3, 3, 1, 3, 3, 1, 3],
        [0, 0, 0, 0, 0, 0, 0],
        [0, 0, 0, 0, 0, 0, 0],
        [0, 0, 0, 0, 0, 0, 0],
    ]
    fills = ["#F3F4F6", "#DCE8F2", "#8EABC3", "#2E5B88"]
    symbols = ["0", "L", "M", "H"]
    image = Image.new("RGB", (1540, 1056), "white")
    draw = ImageDraw.Draw(image)
    draw.text((20, 18), "Google Trends Web Search за пять лет", font=font(40, True), fill="#000000")
    grid_left, grid_top = 405, 130
    cell_w, cell_h = 155, 76
    for col, market in enumerate(["DE", "NL", "BE", "DK", "SE", "FI", "PL"]):
        text_center(draw, (grid_left + col * cell_w, 76, grid_left + (col + 1) * cell_w, grid_top), market, font(28, True), "#111827")
    for row, (family, row_values) in enumerate(zip(families, values)):
        y0 = grid_top + row * cell_h
        y1 = y0 + cell_h
        draw.text((20, y0 + 22), family, font=font(25), fill="#111827")
        for col, value in enumerate(row_values):
            x0 = grid_left + col * cell_w
            x1 = x0 + cell_w
            draw.rectangle((x0, y0, x1, y1), fill=fills[value], outline="#FFFFFF", width=3)
            text_center(draw, (x0, y0, x1, y1), symbols[value], font(27, True), "#FFFFFF" if value >= 2 else "#111827")
    save_chart(image, WORK / "trends_heatmap.png")


make_charts()

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.62)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)
section.different_first_page_header_footer = True

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10.2)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.12

for nm, size in (("Title", 30), ("Subtitle", 14), ("Heading 1", 18), ("Heading 2", 13.5)):
    st = styles[nm]
    st.font.name = "Arial"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(BLACK)
    if nm.startswith("Heading") or nm == "Title":
        st.font.bold = True
    if nm == "Title":
        pPr = st._element.get_or_add_pPr()
        for border in list(pPr.findall(qn("w:pBdr"))):
            pPr.remove(border)

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run("AVDAR MARKET   DEEP RESEARCH   14.09.2026")
set_run_font(hr, size=7.5, color=MID)
footer = section.footer
fp = footer.paragraphs[0]
fr = fp.add_run("Рабочий отчёт для принятия решения   ")
set_run_font(fr, size=7.5, color=MID)
add_page_number(fp)

# Cover
p = doc.add_paragraph(style="Title")
remove_paragraph_borders(p)
p.paragraph_format.space_before = Pt(64)
p.paragraph_format.space_after = Pt(18)
r = p.add_run("Avdar Market исследование спроса и сценариев запуска")
set_run_font(r, size=30, bold=True, color=BLACK)
p = doc.add_paragraph(style="Subtitle")
p.paragraph_format.space_after = Pt(24)
r = p.add_run("Google Trends TikTok Instagram и экономика ассортимента")
set_run_font(r, size=14, color=MID)

add_para(doc, "Срез данных: 14 сентября 2026 года", before=18, after=5)
add_para(doc, "География: Германия, Нидерланды, Бельгия, Дания, Швеция, Финляндия и Польша", after=5)
add_para(doc, "Основа: master-таблица Avdar Market, два исследовательских чата, Google Trends, TikTok, Instagram и Meta Ad Library", after=22)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(28)
p.paragraph_format.space_after = Pt(10)
p.paragraph_format.line_spacing = 1.2
r = p.add_run("Главный вывод")
set_run_font(r, size=13, bold=True, color=BLACK)
add_para(doc, "Рейтинг рынков сохраняется, но выбор первого товара следует считать открытым до физической проверки. Cashmere scarf остаётся лучшим кандидатом для платного DTC-теста; mixed gift bundle усиливает средний чек; пледы подходят для story-led контента и ограниченного запуска; носки имеют самый широкий поисковый и social precedent, но текущая доставка разрушает их экономику как самостоятельного cold-acquisition оффера.", after=8, line=1.2)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(34)
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run("Подготовлено для продолжения исследования и выбора Wave 1")
set_run_font(r, size=9.5, italic=True, color=MID)

page_break(doc)

# Executive summary
add_heading(doc, "Резюме для решения", 1)
add_para(doc, "На 14 сентября 2026 года достаточно данных, чтобы сохранить порядок рынков DE 90, NL 89, BE 81, DK 78, SE 76, PL 73, FI 55. Новые social и Trends слои не меняют этот рейтинг: они уточняют продукт и креатив, а не доказывают более выгодную логистику или конверсию. [S1–S6]", after=7)

canonical_rows = [
    ["Рынки", "DE и NL — Wave 1A; BE и PL — следующая волна; DK и SE условны; FI позже", "Высокая"],
    ["Первый продукт", "Не зафиксирован окончательно. Cashmere scarf — экономический front-runner после QA", "Средняя"],
    ["Bundles", "Mixed scarf gift bundle — слой AOV и gifting, не отдельное доказательство спроса", "Средняя"],
    ["Носки", "Short 3-pack и singles — organic, attach и gifting. Cold paid при текущем freight не проходит", "Высокая"],
    ["Пледы", "Yak и selected wool — content/limited-drop кандидаты; cashmere throw пока hold", "Средняя"],
    ["Каналы", "DTC сначала; OTTO только второй стадией для cashmere; bol и Amazon не Wave 1", "Высокая"],
]
add_table(doc, ["Область", "Каноническое решение", "Уверенность"], canonical_rows, widths=[1.15, 4.9, 0.9], font_size=9.0, caption="Таблица 1  Что теперь считать каноническим")

add_para(doc, "Критический ограничитель. Ни одна товарная гипотеза не готова к платному масштабу, пока не подтверждены физические остатки в Испании, качество образцов, точные размеры упаковки, договорные тарифы прямой и обратной доставки, доказательства состава и прослеживаемости, GPSR и EPR, а также consent, checkout и event tracking. Это не исследовательские мелочи: они определяют допустимые claims, возвраты и CAC ceiling. [S1]", bold_lead="Критический ограничитель.", before=4, after=7)

add_para(doc, "Самое важное противоречие. Instagram и TikTok дают сильный коммерческий precedent для sock packs, но Avdar unit economics отвергает этот вывод как launch decision: short 3-pack даёт только около €2–6 до оплаты и CAC в DE, NL, BE и FI и уходит в минус в DK, SE и PL; long 3-pack отрицателен во всех семи рынках. Social evidence определяет формат креатива, а экономика — право покупать трафик.", bold_lead="Самое важное противоречие.", after=7)

doc.add_picture(str(WORK / "market_scores.png"), width=Inches(6.85))
p = doc.paragraphs[-1]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap = doc.add_paragraph()
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.paragraph_format.space_after = Pt(4)
r = cap.add_run("Рисунок 1  Актуальный scorecard master-таблицы")
set_run_font(r, size=8.5, italic=True, color=MID)

# Scope and method
add_heading(doc, "Объём исследования и метод", 1)
add_para(doc, "В выборку вошли 45 SKU из Product Master. Цвета и размеры не считались отдельными рынками спроса: SKU сведены в 11 поисково различимых семейств по предмету и материалу. Суммарно в клиентской матрице 380 единиц с landed value €6 720,77, но поля Stock Spain, Ordered и ETA не заполнены. Поэтому количество ниже означает плановую матрицу, а не подтверждённый продаваемый остаток. [S1]")

add_para(doc, "Google Trends. Для каждого семейства использовались локальные запросы в семи странах и окна пять лет, 12 месяцев и 90 дней в Web Search. Shopping и YouTube проверялись выборочно: почти все узкие термины оказались ниже разрешения, кроме слабого повторяемого сигнала Wolldecke в DE. Trends показывает относительный индекс 0–100 внутри выбранного сравнения, использует выборку запросов и может показывать ноль при малом объёме; индекс нельзя превращать в абсолютный объём. [S4–S7]", bold_lead="Google Trends.")
add_para(doc, "TikTok. Использованы официальный Creative Center и EU Ads Library, публичные video pages с direct metadata и анонимная поисковая выдача. В шести рынках, кроме DE, актуальный hashtag country filter недоступен. Первые результаты поиска персонализируемы, а keyword search библиотеки не является точным фразовым поиском. [S8–S18]", bold_lead="TikTok.")
add_para(doc, "Instagram и Meta. Organic слой включает публичные профили, последние 9–12 постов, отдельные Reels, видимые реакции и официальные storefront. Meta Ad Library использована как evidence активности, сообщений, таргетинга и длительности. Эти данные не содержат коммерческого ROAS, CAC или продаж; суммарный EU reach по нескольким объявлениям не является уникальной аудиторией. [S19–S31]", bold_lead="Instagram и Meta.")

add_para(doc, "Уровни Trends в матрице: H — устойчивый повторяющийся сигнал; M — примерно 20–49% недель с ненулевым значением; L — примерно 2–19%; 0 — менее 2% или ниже разрешения. Третий символ показывает 90-дневное состояние: ↑ — устойчивое предсезонное усиление; · — ровно или шумно; 0 — ниже разрешения.", before=5, after=7)

page_break(doc)

# Inventory
add_heading(doc, "Ассортимент и экономическая роль", 1)
inventory_rows = [
    ["Плед из яка", "1", "5", "€103,03", "€515,15", "Limited premium drop после GSM и handfeel QA"],
    ["Пледы из овечьей шерсти", "9", "33", "€56,61–98,91", "€2 290,87", "Selected SKU content test; premium price conditional"],
    ["Кашемировые пледы", "2", "6", "€168,72", "€1 012,30", "Hold: €259–288 price floor и фактический pack около 2,9 кг"],
    ["Верблюжья шаль", "1", "7", "€33,37", "€233,58", "Editorial или gift attach после проверки материала"],
    ["Шали из овечьей шерсти", "2", "5", "€35,68", "€178,39", "Небольшой editorial/gift слой"],
    ["Кашемировые шарфы", "6", "49", "€16,06", "€787,05", "Условный hero DTC; stock и replenishment ограничивают scale"],
    ["Шарфы из овечьей шерсти", "3", "28", "€8,13", "€227,70", "Лучший add-on к cashmere scarf"],
    ["Короткие шерстяные носки", "7", "113", "€4,50", "€508,50", "3-pack, add-on, gift; не cold acquisition"],
    ["Длинные шерстяные носки", "8", "121", "€6,93", "€838,53", "Attach only: 3-pack отрицателен во всех рынках"],
    ["Носки из яка", "2", "0", "€4,50", "€0", "Blocked: нет количества и wear evidence"],
    ["Шерстяные перчатки", "4", "13", "€9,90", "€128,70", "Gift/scarf attach после fit и seam test"],
]
add_table(doc, ["Семейство", "SKU", "Ед.", "Landed за ед.", "Landed value", "Роль"], inventory_rows, widths=[1.55,0.42,0.45,0.9,0.92,2.55], font_size=8.0, caption="Таблица 2  Все товарно материальные семейства")

doc.add_picture(str(WORK / "inventory.png"), width=Inches(6.85))
p = doc.paragraphs[-1]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap = doc.add_paragraph()
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = cap.add_run("Рисунок 2  Распределение матрицы по единицам и landed value")
set_run_font(r, size=8.5, italic=True, color=MID)

add_para(doc, "Пледы занимают 56,8% landed value, а шерстяные носки — 61,6% всех единиц. Поэтому первая группа несёт риск замороженного капитала, вторая — риск большого остатка в товаре с низким ticket и отрицательной доставкой. Cashmere scarf, напротив, содержит только 11,7% landed value и даёт лучший запас contribution, но 49 единиц недостаточно для двух полноценных country gates без пополнения.", after=4)

# Trends
add_heading(doc, "Google Trends по всем товарным семействам", 1)
add_para(doc, "Пять лет показывают два повторяющихся cross-market двигателя: wool socks и sheep-wool blankets. Cashmere scarf имеет сезонный, но воспроизводимый спрос в DE и PL, пригодный для теста сигнал в DK и более слабый слой в NL. Exact yak, camel и gift phrases почти всегда ниже разрешения. Это означает, что редкий материал следует продавать как доказанный атрибут внутри широкой категории, а не строить acquisition только на exact material keyword. [S4–S7]")

doc.add_picture(str(WORK / "trends_heatmap.png"), width=Inches(6.75))
p = doc.paragraphs[-1]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap = doc.add_paragraph()
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = cap.add_run("Рисунок 3  Уровень повторяемости Web Search за пять лет")
set_run_font(r, size=8.5, italic=True, color=MID)

trend_rows1 = [
    ["Cashmere scarf", "M/M/↑", "L/L/·", "0/L/·", "M/M/·"],
    ["Wool scarf", "H/H/↑", "M/M/·", "L/L/·", "L/L/·"],
    ["Wool socks", "H/H/↑", "H/H/↑", "M/M/·", "M/H/↑"],
    ["Yak socks", "0/0/0", "0/0/0", "0/0/0", "0/L/0"],
    ["Wool gloves", "M/M/0", "L/L/·", "0/0/·", "0/0/0"],
    ["Camel shawl", "0/0/0", "0/0/0", "0/0/0", "0/0/0"],
    ["Wool shawl", "L/M/·", "0/L/·", "0/0/·", "0/0/·"],
    ["Sheep-wool throw", "H/H/↑", "H/H/↑", "L/L/·", "H/H/↑"],
    ["Yak throw", "0/0/0", "0/0/0", "0/0/0", "0/0/0"],
    ["Cashmere throw", "0/L/·", "0/L/0", "0/0/·", "0/L/·"],
    ["Gift exact", "0/0/0", "0/0/0", "0/0/0", "0/0/0"],
]
add_table(doc, ["Семейство", "DE", "NL", "BE", "DK"], trend_rows1, widths=[2.35,1.0,1.0,1.0,1.0], font_size=8.5, caption="Таблица 3  Trends пять лет 12 месяцев 90 дней часть 1")
trend_rows2 = [
    ["Cashmere scarf", "L/L/0*", "0/0/0", "M/M/↑"],
    ["Wool scarf", "L/L/0*", "L/L/·", "M/M/↑"],
    ["Wool socks", "H/H/↑", "H/H/↑", "M/M/↑"],
    ["Yak socks", "0/0/0", "0/0/0", "0/L/0"],
    ["Wool gloves", "0/0/·", "0/0/0", "M/M/·"],
    ["Camel shawl", "0/0/0", "0/0/0", "0/0/0"],
    ["Wool shawl", "0/0/·", "0/0/·", "L/M/·"],
    ["Sheep-wool throw", "H/H/↑", "L/L/↑", "H/H/↑"],
    ["Yak throw", "0/0/0", "0/0/0", "0/0/0"],
    ["Cashmere throw", "0/0/0", "0/0/·", "0/0/·"],
    ["Gift exact", "0/0/0", "0/0/0", "0/0/0"],
]
add_table(doc, ["Семейство", "SE", "FI", "PL"], trend_rows2, widths=[2.55,1.25,1.25,1.25], font_size=8.5, caption="Таблица 4  Trends пять лет 12 месяцев 90 дней часть 2")
add_para(doc, "* Шведские слитные exact compounds систематически недосчитывают: варианты с пробелом и альтернативным написанием повышают долю ненулевых недель, но cashmere и wool scarves всё равно остаются long-tail. Бельгийские blanket terms также дробятся между Dutch и French; поэтому BE row имеет более низкую уверенность.", before=0, after=6)

support_rows = [
    ["DE", "37,8%", "92,4%", "Wool scarf 56,9%; gloves 22,5%", "Высокая по core"],
    ["NL", "8,8%", "75,5% за 12m", "Wool throw 97,7%", "Высокая по throws"],
    ["BE", "1,9%", "20,2%", "Blanket phrase 2,7%", "Средняя или низкая"],
    ["DK", "21,0%", "42,0%", "Wool throw 77,1%", "Высокая по core"],
    ["SE", "8–10% variants", "63,7%", "Wool throw 93,5%", "Низкая по exact premium"],
    ["FI", "0% exact", "100%", "Wool throw 12,6%", "Высокая по socks"],
    ["PL", "29,8%", "44,7%", "Wool throw 73,7%; gloves 21,4%", "Высокая"],
]
add_table(doc, ["Рынок", "Cashmere scarf", "Wool socks", "Другой 5y сигнал", "Confidence"], support_rows, widths=[0.55,1.25,1.25,2.45,1.5], font_size=8.2, caption="Таблица 5  Доля пятилетних недель с индексом выше нуля")
add_para(doc, "Доли недель используются только как проверка повторяемости, а не как market size. Для NL socks приведён годовой ряд из-за нестабильной dynamic table. [S35–S43]")

add_heading(doc, "Сезонность и расхождения", 2)
add_bullets(doc, [
    ("Пять лет. ", "Scarves, blankets и socks устойчиво концентрируются в Q4. Wool socks и wool blankets остаются заметными и вне пика."),
    ("Последние 12 месяцев. ", "Порядок не изменился: структурного premium-breakthrough не видно. DK cashmere стал заметнее, NL остаётся умеренным, FI exact cashmere почти отсутствует."),
    ("Последние 90 дней. ", "С конца августа и в сентябре усиливаются core wool и DE/PL scarf terms. Однодневные spikes редких запросов исключены как sampling noise."),
    ("Shopping и YouTube. ", "Узкие material-product terms в основном ниже разрешения. Их нельзя использовать как самостоятельное подтверждение paid-shopping или video demand."),
])

contradictions = [
    ["SE kashmirhalsduk", "HYPD 590 в месяц", "0,8% недель exact; 8–10% с variants", "Не использовать HYPD как exact volume"],
    ["FI kashmirhuivi", "HYPD 260 в месяц", "Ноль exact; variants около 0,4–0,8%", "Самое сильное расхождение"],
    ["DE Yak Socken", "HYPD 210 в месяц", "1,1% недель за пять лет; ноль за 12 месяцев", "Только experimental long-tail"],
    ["NL cashmere", "HYPD 880 в месяц", "8,8% недель за пять лет; 17% за 12 месяцев", "Сезонность совпадает, scale не подтверждён"],
    ["BE cashmere", "HYPD NL 170 + FR 260", "Около 1,9% недель за пять лет", "Разделять FR и NL и не суммировать механически"],
    ["FI villahuopa", "HYPD 1 000 в месяц", "12,6% недель за пять лет", "Moderate mismatch"],
]
add_table(doc, ["Термин", "Absolute layer", "Google Trends", "Решение"], contradictions, widths=[1.35,1.35,2.25,2.05], font_size=8.1, caption="Таблица 6  Где абсолютные оценки не согласуются с Trends")

add_para(doc, "Rising и geography. В DE усиливались modifiers для women, black, colourful и warm; в NL — beige, alpaca и merino; в BE — gender и `100% cachemire`; в PL — merino, men, women, warmth и размеры. Breakout означает рост более 5 000%, а не большой абсолютный рынок. Cashmere чаще концентрировался в capital and urban areas, тогда как wool socks были относительно сильнее в холодных северных и периферийных регионах. Эти индексы показывают долю интереса внутри региона, а не число поисков. [S44–S45]", after=6)

add_para(doc, "Практическое следствие. Search acquisition следует строить на category terms с доказанной повторяемостью, а `yak`, `camel`, `100%`, origin и craft использовать в ad copy, PDP и контенте как reason to believe. Exact rare-fiber terms подходят для SEO-страниц, remarketing и discovery, но не для планирования объёма продаж.", bold_lead="Практическое следствие.", after=5)

# TikTok
add_heading(doc, "TikTok как канал создания спроса", 1)
add_para(doc, "TikTok лучше всего подходит Avdar для объяснения материала, стилизации и снятия возражений. Natural-fiber terms не входят в три анонимно доступных top hashtags DE, но в 30-дневном Apparel and Accessories слое присутствуют `herbstoutfit` и `autumnoutfit`. Это поддерживает pre-fall scarf framing, но не даёт объёма продаж. [S8–S9]")

tiktok_rows = [
    ["Cashmere impulse", "643 тыс.", "60,3 тыс.", "Желание и вопрос о покупке", "S10"],
    ["OFT Germany cashmere scarf", "316 тыс.", "7 264", "Локальный outfit и styling", "S11"],
    ["Falconeri Pre-Fall OOTD", "2,1 млн", "10,2 тыс.", "Брендовый commercial OOTD", "S12"],
    ["Cashmere и wool care", "216,6 тыс.", "12,5 тыс.", "Evergreen tutorial", "S13"],
    ["Wool blanket production", "569,5 тыс.", "5 710", "Процесс сильнее catalogue shot", "S14"],
    ["Nordic blanket ritual", "12,8 млн", "444,1 тыс.", "Культурный editorial reach", "S15"],
    ["Why yak is great", "29,3 тыс.", "4 695", "Education и discussion", "S16"],
    ["PL merino socks около 17 zł", "201,5 тыс.", "5 388", "Commodity price pressure", "S17"],
]
add_table(doc, ["Публичный сигнал", "Views", "Likes", "Что проверяет", "Источник"], tiktok_rows, widths=[2.15,0.82,0.82,2.45,0.55], font_size=8.2, caption="Таблица 7  Репрезентативные TikTok видео на дату среза")
add_para(doc, "Эти публикации не являются контролируемым экспериментом. В частности, 12,8 млн просмотров культурного ритуала с пледом не означает готовность купить Avdar throw, а 2,1 млн у Falconeri не раскрывают вклад бренда, media spend или conversion. Значение выборки — в creative territory и качестве hook.", after=6)

market_tiktok = [
    ["DE", "Cashmere scarf и outfit; wool blanket уходит в DIY", "Scarf hero; socks и throws через problem/process"],
    ["NL", "Cashmere signal умеренный; `wool socks` collision с fintech", "Материал + предмет + use-case в spoken hook и overlay"],
    ["BE", "French scarf UGC мал; wool education заметнее", "Разделять FR и NL; education важнее polished catalogue"],
    ["DK", "Низкая saturation; Uldplaiden владеет exact, но reach мал", "Hygge, cold use-case, provenance и ASMR"],
    ["SE", "Cashmere collides с Kashmir; creator integration даёт reach", "Локализовать `ullstrumpor`; creator/use-case вместо flat lay"],
    ["FI", "Wool socks — craft и Christmas culture; compounds ломаются", "Тестировать варианты написания и gifting heritage"],
    ["PL", "Самая плотная коммерческая выдача и price anchors", "Не конкурировать скидкой; доказывать состав и durability"],
]
add_table(doc, ["Рынок", "Что видно", "Решение для Avdar"], market_tiktok, widths=[0.55,3.05,3.20], font_size=8.4, caption="Таблица 8  Market readout TikTok")

add_heading(doc, "Рабочие форматы", 2)
add_bullets(doc, [
    "Search-native tutorial: четыре способа завязать шарф, один шарф в пяти образах, ответ на комментарий.",
    "Quality interrogation: почему похожие cashmere products стоят по-разному; состав, fibre length, вес, finish и blend comparison.",
    "Care utility: стирка, удаление катышков, хранение и проветривание.",
    "Problem and solution: холодные ноги, тепло без объёма, layering в поездке.",
    "Process and provenance: прядение, ткачество, руки мастера и origin. Claims допускаются только после supplier proof.",
    "Sensory and choice: macro texture, драпировка, ASMR, выбор цвета и получателя.",
    "Gift ritual: распаковка от лица получателя, персональная записка и last-order date.",
])

add_para(doc, "В TikTok EU Ads Library найден один свежий cashmere advertiser с Sales objective и broad-Europe targeting всех семи рынков; видимый диапазон unique users — 0–1 тыс. на рынок. Это подтверждает наличие paid-конкурента, но не scale или sales. Exact brand queries по основным референсам дали ноль в заданном DE окне, что нельзя интерпретировать как отсутствие рекламы вообще. [S18]", after=5)

# Instagram and Meta
add_heading(doc, "Instagram и Meta как слой доверия и визуального доказательства", 1)
add_para(doc, "В Instagram повторяются шесть устойчивых territory: происхождение и ремесло; tactile slow comfort; styling; material education; multi-buy; gift-by-recipient. Reels нужны для движения, styling, offer и unboxing, но static и carousel остаются рабочими форматами для ассортимента, фактуры и образования. [S19–S27]")

ig_rows = [
    ["GOBI", "Heritage, herders, fibre to product", "Global verified account; локальные DE/PL аккаунты переводят аудиторию в global", "Origin system и bundles"],
    ["PURSCHOEN", "Street styling, vivid colour, creator layer", "5–7 posts в неделю; scarf-native", "Главный scarf creative benchmark"],
    ["Falconeri", "Premium sensory ritual", "Polished static/carousel, почти ежедневно", "Art direction и tactile language"],
    ["ROECKL", "Heritage, recipient, gift guide", "Gift and texture combinations", "Gift box должен решать occasion"],
    ["Nordic Socks", "3- и 5-pair packs, comfort, gifting", "169 тыс. followers, но слабые recent visible likes", "Pack architecture, не sales proof"],
    ["Uldplaiden", "2+ −20%, free-shipping threshold", "Reels-heavy и local Danish", "Offer repetition и urgency"],
    ["Mongolian.pl", "Long-form material education", "Local Polish static/carousels", "Снимать price objection через свойства"],
    ["Balmuir и Silkeborg", "Home ritual, colour, Nordic slow living", "Static/carousel и provenance", "Throw storytelling"],
]
add_table(doc, ["Референс", "Creative territory", "Наблюдение", "Применение"], ig_rows, widths=[1.15,2.05,2.25,1.55], font_size=8.1, caption="Таблица 9  Organic Instagram benchmarks")

add_para(doc, "GOBI Polska 13 сентября 2026 года пригласил аудитории локальных DE, PL, UK, US и FR аккаунтов переходить в единое global community. Для Avdar это сильный operational precedent: начинать с одного основного аккаунта и локализованных captions, overlays, creators и landing pages; отдельные country handles создавать только при достаточном локальном cadence. [S19]", after=6)

meta_rows = [
    ["GOBI Cashmere", "48 из 50", "14", "Craft, Mongolian origin, wardrobe; 45% anniversary offer", "All-gender EU; € spend не раскрыт"],
    ["ROECKL", "36 из 40", "29", "Handmade accessories, autumn looks, cashmere and wool", "Сильная DE seasonal burst"],
    ["Uldplaiden", "27", "1 на дату среза", "2+ −20%, home/outdoor use, corporate gifting", "25 из 27 targeted women"],
    ["Dinadi", "2", "2", "Undyed yak, rarity, 40 hours hand-knit", "Women; provenance-first"],
    ["Joe Merino", "40", "26", "Colour drops, premium blend, broader wardrobe", "Полезен как apparel system, не sock demand"],
]
add_table(doc, ["Advertiser", "Релевантные rows", "Active", "Сообщение", "Ограничение"], meta_rows, widths=[1.25,1.0,0.8,2.55,1.45], font_size=8.0, caption="Таблица 10  Meta Ad Library capture 14 сентября 2026 года")

add_para(doc, "У Uldplaiden суммарная недедуплицированная delivery по полученным EU rows на 85% приходилась на возраст 55+. Это не профиль покупателя и не уникальная аудитория, но достаточно сильный сигнал, чтобы не строить throw strategy только вокруг молодой TikTok-аудитории. Для пледов Instagram Feed и Reels с mature female creative deserve отдельный тест. [S28–S31]", after=6)
add_para(doc, "Generic keyword searches сильно загрязнены. В выдаче `cashmere scarf` 24 из 50 строк относились к short-drama advertisers, а 21 — к одному scarf advertiser. `cashmere blanket` почти полностью состоял из нерелевантных сюжетных объявлений. Поэтому counts без ручной advertiser-level очистки не являются размером категории.", after=5)

# Launch models
add_heading(doc, "Сравнение пяти моделей запуска", 1)
model_rows = [
    ["1  Cashmere scarf DTC", "Поиск M в DE/PL; сильный TikTok/IG", "Лучший запас contribution; CAC ceiling около €45–47 после full-stack guardrails", "49 ед.; QA, claims, pack, replenishment", "Приоритет 1 условно"],
    ["2  Mixed gift bundle", "Сильный Q4 precedent и unboxing", "Wool scarf add-on даёт около €23–51 incremental contribution; DE attach repair около 6–7%", "1+1 pack, partial return, 28 wool scarves", "Запускать рядом с hero"],
    ["3  Yak и selected wool throws", "Wool throw H; high story fit", "Yak @€219 даёт около €41–62 pre-fee/CAC; wool зависит от cohort", "44 ед. total; GSM, handfeel, pilling, bulky reverse", "Organic и limited drop"],
    ["4  Short sock 3-pack", "Самый широкий category demand; strong pack precedent", "Около €2–6 pre-fee/CAC только в DE/NL/BE/FI; минус в DK/SE/PL", "Sizing, reinforcement, 3-pack tariff", "Attach/gift; не cold paid"],
    ["5  Single socks", "Понятный entry и utility creative", "Структурно отрицательны с merchant-paid Spain freight", "Размеры и returns; low ticket", "Add-on, sampling, seeding"],
]
add_table(doc, ["Модель", "Demand и social", "Экономика", "Hard gate", "Решение"], model_rows, widths=[1.15,1.25,2.0,1.55,1.05], font_size=7.7, caption="Таблица 11  Итоговая модель выбора")

add_heading(doc, "Почему cashmere scarf остаётся первым", 2)
add_para(doc, "При €129 cashmere scarf выдерживает VAT, landed cost и Spain outbound лучше остальных самостоятельных офферов. В предыдущей модели DE и NL оставляли порядка €45–47 допустимого CAC после полного набора guardrails. Search в DE и PL повторяем, в NL сезонно достоверен; TikTok и Instagram дают несколько рабочих форматов, а exact-format comparator prices подтверждают premium lane. Ограничение не в отсутствии creative idea, а в неподтверждённом качестве, происхождении, pack и replenishment. [S1, S10–S13, S19–S21]")

add_heading(doc, "Почему bundle является вторым слоем", 2)
add_para(doc, "Cashmere scarf плюс wool scarf при скидке 10–15% на второй товар даёт наиболее убедительный AOV repair: около €23–51 ожидаемой incremental contribution по рынкам, если два изделия остаются в том же parcel tier. Социальные референсы подтверждают recipient, colour choice, gift note и set-builder механику. Но bundle нельзя считать доказанным, пока не измерены закрытая упаковка, тариф, частичный возврат и attach conversion. [S1, S22–S26, S34]")

add_heading(doc, "Почему socks не становятся первым paid hero", 2)
add_para(doc, "И Trends, и social показывают широкую wool-sock category и понятный 3-pair precedent. Однако при существующей доставке contribution слишком мала даже до payment fee, returns и CAC. Тестировать packs нужно органически, в gift layer и как add-on. Paid acquisition разрешается только после изменения price, pack или shipping, которое создаёт достаточный contribution guardrail. Yak socks дополнительно заблокированы нулевым количеством и отсутствием wear and wash proof.")

add_heading(doc, "Почему throws идут отдельной дорожкой", 2)
add_para(doc, "Throws дают самый сильный process, provenance и home-ritual контент, а generic wool blanket спрос устойчив. Но cashmere throws требуют примерно €259–288 только для около €30 contribution до fee, returns и CAC; sheep-wool cohorts требуют разной цены €128–200; yak throw имеет всего пять единиц и nominal GSM около 282. Поэтому сначала снимаются доказательства и organic intent, затем возможен небольшой limited drop без обещания scale.")

# Channel and markets
add_heading(doc, "Рынки и каналы", 1)
market_rows = [
    ["DE 90", "Wave 1A", "Cashmere scarf; bundle", "Google Search/Shopping после QA; Instagram styling; TikTok tutorial", "Главный риск — proof, service и promo crowding"],
    ["NL 89", "Wave 1A", "Cashmere scarf; bundle", "DTC, Dutch PDP; IG central account; TikTok term disambiguation", "Textile EPR и exact freight"],
    ["BE 81", "Wave 1B", "Scarf и bundle", "Отдельные NL и FR cells", "Не суммировать bilingual demand"],
    ["DK 78", "Wave 1B условно", "Scarf; throw organic", "Local offer creative и mature-female Meta test", "Freight score 2 из 5"],
    ["SE 76", "Hold или premium reopen", "Throw story; scarf test later", "Instagram home/lifestyle; local compounds", "Trends и absolute volume расходятся"],
    ["PL 73", "Wave 1B", "Scarf; sock education", "Polish material education и low-CPC search", "Жёсткие sock price anchors"],
    ["FI 55", "Later", "Wool socks organic; throw lifestyle", "English master creative + Finnish commerce", "Payment, freight и слабый exact cashmere"],
]
add_table(doc, ["Рынок", "Роль", "Продукт", "Канал", "Ограничение"], market_rows, widths=[0.75,1.0,1.25,2.65,1.35], font_size=7.9, caption="Таблица 12  Market channel allocation")

channel_rows = [
    ["DTC", "Первый этап", "Контроль цены, story, PDP, bundles и данных", "Нужны QA, tracking, returns и local commerce"],
    ["OTTO DE", "Второй этап только cashmere", "При €119 около €47,09 до returns, ads и monthly fee", "€99,90 в месяц; 15% Accessories; material proof"],
    ["bol NL/BE", "Позже", "Доступ к локальному demand", "Invite-only, ≥€500 тыс. first-year potential, Dutch service"],
    ["Amazon DE socks", "Не Wave 1", "Большой generic shelf", "Низкая средняя цена, 90% sponsored, тысячи reviews"],
    ["Marketplaces для throws", "Hold", "Возможен demand capture", "Bulky shipping, returns, fee и claims уменьшают margin"],
]
add_table(doc, ["Канал", "Статус", "Плюс", "Почему не сейчас"], channel_rows, widths=[1.25,1.35,2.25,2.15], font_size=8.4, caption="Таблица 13  DTC и marketplaces")

add_para(doc, "Account architecture. На старте рационален один global Instagram и один TikTok account с локальными spoken hooks, overlays, captions и destination pages. DE и NL paid campaigns остаются раздельными. BE NL и FR ведутся как два самостоятельных cells. Country accounts появятся только тогда, когда есть контент-план и community management на языке рынка.", bold_lead="Account architecture.", after=5)

# Data gaps
add_heading(doc, "Приоритетный план закрытия Data Gaps", 1)
add_para(doc, "Исходный Data Gaps содержит 38 строк. Ниже они сведены в последовательность, которая либо разрешает, либо запрещает трафик. Работа по low-priority benchmark не должна опережать физический товар, shipping и compliance. [S1]")

p0_rows = [
    ["1", "Физический stock Spain по SKU, reserve, Ordered и ETA", "Клиент", "Блокирует выбор hero и объём теста"],
    ["2", "Cashmere scarf QA и два representative colours", "Клиент и лаборатория", "Состав, GSM, handfeel, pilling, colour transfer, fringe"],
    ["3", "Socks и gloves fit, reinforcement, seam, wash and abrasion", "Клиент", "Блокирует pack claims и returns"],
    ["4", "Throw GSM, finish, shedding, shrinkage и use-case", "Клиент", "Разделяет editorial plaid и premium blanket"],
    ["5", "Пакеты 1 item, 1+1, 3 socks и throw", "Fulfillment", "Записать closed dimensions и gross weight"],
    ["6", "Contracted outbound и reverse quotes DE/NL", "Fulfillment", "Одинаковые pack и service definitions"],
    ["7", "Supplier claim dossier, batch, GPSR, REACH и EPR", "Supplier и legal", "Запрещает неподтверждённые origin, micron, no-itch claims"],
    ["8", "NL textile EPR route и EU guarantee notice", "Legal и client", "Notice должен быть внедрён к 27.09.2026 [S32–S33]"],
]
add_table(doc, ["№", "До любого paid traffic", "Владелец", "Acceptance evidence"], p0_rows, widths=[0.3,3.3,1.2,2.2], font_size=8.2, caption="Таблица 14  Gate 0")

p1_rows = [
    ["Returns", "Локальная policy, cancellation until fulfilment, partial bundle return, reverse tariff"],
    ["Checkout", "Payment methods, 3DS, delivery promise, confirmation emails и refund path"],
    ["Measurement", "Consent Mode, Shopify Customer Privacy, Meta Pixel and CAPI, UTMs, purchase order ID"],
    ["Geo и language", "Presence targeting, separate market campaigns, BE NL и FR split"],
    ["Merchant Center", "Shipping, returns, weight, currency и diagnostics после финального pack"],
    ["Price governance", "Стабильный hero price и архив offer state; не смешивать тест с promo regime"],
]
add_table(doc, ["Блок", "Что должно быть доказано до первого cohort"], p1_rows, widths=[1.2,5.8], font_size=8.8, caption="Таблица 15  Gate 1")

add_heading(doc, "Что можно отложить", 2)
add_para(doc, "Дополнительные pure-camel benchmarks в BE и SE, расширение proprietary trade datasets, broad Meta keyword counts, bol commission hunting и отдельные country social accounts не блокируют Wave 1. Их следует открывать только если физический продукт и первые cohorts делают решение зависимым от этих данных.")

# Experiment plan
add_heading(doc, "План экспериментов на 90 дней", 1)
timeline_rows = [
    ["14–21 сентября", "Gate 0 physical sprint", "Inventory reconciliation, representative QA, pack measurements, claims dossier, shipping RFQ", "Go or no-go по cashmere scarf, bundle и selected throw"],
    ["22 сентября–5 октября", "Storefront и content build", "DE/NL PDP, EU notice, returns, tracking QA, 18–24 creative assets", "Каждый asset имеет product, hook, market и destination ID"],
    ["6–19 октября", "Organic smoke test", "Scarf styling/proof, bundle recipient, throw process, socks problem/solution", "Ранжирование по hold, completion, saves, shares и qualified PDP visits"],
    ["20 октября–15 ноября", "DTC Wave 1A", "DE и NL раздельно; stable price; Search/Shopping и selected social creatives", "DE 10 fulfilled ≤255 clicks; NL 10 fulfilled ≤386 clicks"],
    ["16 ноября–15 декабря", "Maturation и Q4 decision", "Returns, payment leakage, reverse cost, defects, bundle attach", "Scale только после matured contribution; иначе iterate или stop"],
]
add_table(doc, ["Период", "Этап", "Работа", "Выход"], timeline_rows, widths=[1.0,1.2,2.8,2.0], font_size=8.0, caption="Таблица 16  Исполняемая последовательность")

add_heading(doc, "Минимальная creative test matrix", 2)
creative_rows = [
    ["Cashmere scarf", "4 ways to wear", "Styling tutorial", "DE и NL", "TikTok и IG Reels"],
    ["Cashmere scarf", "110 g и тепло без объёма", "Product proof после weight QA", "DE и NL", "Reels и carousel"],
    ["Cashmere scarf", "100% versus blend", "Material education", "DE и PL", "TikTok и IG carousel"],
    ["Mixed bundle", "Gift for the person who is always cold", "Recipient and reveal", "DE и NL", "TikTok и IG Reels"],
    ["Mixed bundle", "Choose two colours", "Choice and self-gifting", "DE и NL", "Reels и carousel"],
    ["Short socks", "Cold feet in boots", "Problem and solution", "DE, NL и PL", "TikTok organic"],
    ["Short socks", "Three recipients or three days", "Pack utility", "DE и NL", "TikTok and IG organic"],
    ["Selected throw", "From fibre to loom", "Process and provenance", "DE, DK и SE", "TikTok and IG"],
    ["Selected throw", "Actual size and room before after", "Scale and interior", "DE и NL", "IG Reels and carousel"],
]
add_table(doc, ["Продукт", "Hook", "Territory", "Рынок", "Формат"], creative_rows, widths=[1.1,1.7,1.5,1.0,1.6], font_size=8.1, caption="Таблица 17  Первая партия креативов")

add_para(doc, "Metrics. Для каждого creative ID фиксируются 3- и 6-second hold, completion, saves, shares, comments, profile-to-site CTR, PDP engagement, add to cart, checkout initiation, purchase и matured contribution. Views и likes используются для диагностики hook, но не для product-market decision.", bold_lead="Metrics.")
add_para(doc, "Scale rule. Десять fulfilled orders — статистический acquisition gate, а не автоматическое разрешение увеличить бюджет. Сначала ранний cohort проходит применимое return window; затем учитываются refunds, payment leakage, reverse cost и disposition. Bundle дополнительно должен подтвердить attach rate около 6–7% или выше в DE при неизменном parcel tier. [S1]", bold_lead="Scale rule.", after=6)

add_heading(doc, "Решение после Gate 0", 2)
add_numbered(doc, [
    "Если cashmere scarf проходит stock, QA, pack, claims и replenishment, запускать DTC DE и NL с mixed bundle рядом с hero.",
    "Если scarf не проходит, но selected yak или lower-cost wool throw проходит quality и shipping, делать limited editorial drop и organic demand test без обещания paid scale.",
    "Если оба направления не проходят, не покупать cold traffic. Использовать период для supply correction, supplier evidence и organic learning.",
])

# Canonical updates
add_heading(doc, "Предлагаемые изменения в master таблице", 1)
add_para(doc, "Master-таблица не изменялась в ходе этого deep-dive. После подтверждения владельцем рекомендуется внести новую версию одним append-only блоком, чтобы не размыть историю решений.")
updates = [
    ["Scorecard", "Обновить заголовок рабочей версии до v4.07, сохранив ссылки на v4.04–v4.06 как историю и не меняя ranking"],
    ["Product decision", "Заменить implicit scarf-first на conditional two-track: scarf plus bundle, throws organic lane"],
    ["Keywords", "Добавить Trends directionality и явные contradictions SE, FI, DE yak, NL и BE"],
    ["Competitors", "Добавить TikTok/Instagram benchmark set и статус dormant profiles"],
    ["Data Gaps", "Разделить social public evidence resolved и first-party content performance open"],
    ["Evidence Log", "Добавить source rows этого отчёта с датой 2026-09-14"],
    ["Quality control", "Исправить date-typed numeric cells 08 Shipping G16:H16 и 10 Scorecard J17"],
]
add_table(doc, ["Раздел", "Изменение"], updates, widths=[1.2,5.8], font_size=8.8, caption="Таблица 18  Draft change set")
add_para(doc, "Причина не редактировать master автоматически — новые social и Trends данные содержат разные confidence levels и не должны перезаписывать экономику. Их следует добавить как отдельный evidence layer с датой, источником, методом и decision impact.")

# Sources
add_heading(doc, "Источники", 1)
add_para(doc, "Все web и social источники проверены на дату среза. Динамические счётчики и активность могут измениться. Тексты рекламных объявлений в отчёте пересказаны; Meta archive IDs и direct video links служат точками проверки.")

sources = [
    ("S1", "Исследование Avdar Market master Google Sheet", "https://docs.google.com/spreadsheets/d/1YmtIAePP6YoGmcc-30R37Vyej0Wt9YTMzaDNA_7Nm0Q/edit"),
    ("S2", "Первый исходный исследовательский чат", "https://chatgpt.com/share/6aa7c136-4930-83eb-9b1e-e9c5075a83bf"),
    ("S3", "Второй исходный исследовательский чат", "https://chatgpt.com/share/6aa7c157-b7d4-83ed-9e10-2ef2ec8b441b"),
    ("S4", "Google Trends Basics", "https://newsinitiative.withgoogle.com/resources/trainings/google-trends/basics-of-google-trends/"),
    ("S5", "Google Trends FAQ about data", "https://support.google.com/trends/answer/4365533?hl=en"),
    ("S6", "Advanced Google Trends", "https://newsinitiative.withgoogle.com/resources/trainings/google-trends/advanced-google-trends/"),
    ("S7", "Google Trends API alpha", "https://developers.google.com/search/apis/trends"),
    ("S8", "TikTok official About Creative Center", "https://ads.tiktok.com/resources/help/article/creative-center?lang=en-GB"),
    ("S9", "TikTok official How to use Trends", "https://ads.tiktok.com/resources/help/article/how-to-use-trends?lang=en"),
    ("S10", "TikTok cashmere impulse consumer video", "https://www.tiktok.com/@tavbeans/video/7675058930594172193"),
    ("S11", "TikTok OFT Germany cashmere scarf styling", "https://www.tiktok.com/@oft.st/video/7674187190888418592"),
    ("S12", "TikTok Falconeri Pre Fall OOTD", "https://www.tiktok.com/@falconeri/video/7682887809702563104"),
    ("S13", "TikTok cashmere and wool care tutorial", "https://www.tiktok.com/@catherine_lockhart/video/7575326671155318071"),
    ("S14", "TikTok wool blanket production process", "https://www.tiktok.com/@dweetlive/video/7631590790917115158"),
    ("S15", "TikTok Nordic blanket ritual", "https://www.tiktok.com/@kyanasue/video/7611997731170700566"),
    ("S16", "TikTok yak material education", "https://www.tiktok.com/@wolfvsgoat/video/7593496649381530910"),
    ("S17", "TikTok Poland merino sock price anchor", "https://www.tiktok.com/@promowariatki/video/7557785932993236246"),
    ("S18", "TikTok EU Ads Library cashmere ad 1875987438139633", "https://library.tiktok.com/ads/detail/?ad_id=1875987438139633"),
    ("S19", "Instagram GOBI global profile", "https://www.instagram.com/gobicashmere/"),
    ("S20", "Instagram GOBI cashmere scarf carousel", "https://www.instagram.com/gobicashmere/p/DcvnKnsD8f6/"),
    ("S21", "Instagram PURSCHOEN scarf styling Reel", "https://www.instagram.com/purschoen/reel/DdHKxhPlfGU/"),
    ("S22", "Instagram Nordic Socks 3 pair pack post", "https://www.instagram.com/nordicsocks/p/DdOwhPxiUkG/"),
    ("S23", "Instagram Uldplaiden multi buy Reel", "https://www.instagram.com/uldplaiden/reel/DdO5I9TDNSC/"),
    ("S24", "Instagram Mongolian.pl material education", "https://www.instagram.com/mongolian.sklep/p/DdGIiIQjGax/"),
    ("S25", "Instagram Balmuir slow living Reel", "https://www.instagram.com/balmuir/reel/DclodRYhugm/"),
    ("S26", "ROECKL official gift guide", "https://roeckl.com/de/geschenke-guide"),
    ("S27", "Instagram Silkeborg Uldspinderi colour and throw story", "https://www.instagram.com/silkeborg_uldspinderi/p/DcveuT-o9GD/"),
    ("S28", "Meta official Ad Library API documentation", "https://www.facebook.com/ads/library/api/"),
    ("S29", "Meta Ad Library GOBI Cashmere search", "https://www.facebook.com/ads/library/?active_status=all&ad_type=all&country=DE&q=GOBI%20Cashmere&search_type=keyword_unordered"),
    ("S30", "Meta Ad Library ROECKL search", "https://www.facebook.com/ads/library/?active_status=all&ad_type=all&country=DE&q=ROECKL&search_type=keyword_unordered"),
    ("S31", "Meta Ad Library Uldplaiden search", "https://www.facebook.com/ads/library/?active_status=all&ad_type=all&country=DK&q=Uldplaiden&search_type=keyword_unordered"),
    ("S32", "EU Implementing Regulation 2025 1960", "https://eur-lex.europa.eu/eli/reg_impl/2025/1960/oj"),
    ("S33", "European Commission practical guarantee notice files", "https://commission.europa.eu/publications/practical-guidelines-and-high-resolution-vector-files-eu-notice-and-label-product-guarantees_en"),
    ("S34", "GOBI Cashmere Home Essentials Set", "https://www.gobicashmere.com/products/cashmere-home-essentials-set"),
    ("S35", "Google Trends DE Web Search five year core comparison", "https://trends.google.com/trends/explore?hl=en&date=today%205-y&geo=DE&q=Kaschmirschal%2CWollschal%2CWollsocken%2CWollhandschuhe%2CWolldecke"),
    ("S36", "Google Trends NL Web Search five year core comparison", "https://trends.google.com/trends/explore?hl=en&date=today%205-y&geo=NL&q=kasjmier%20sjaal%2Cwollen%20sjaal%2Cwollen%20sokken%2Cwollen%20handschoenen%2Cwollen%20plaid%20%2B%20wollen%20deken"),
    ("S37", "Google Trends BE Web Search five year bilingual core comparison", "https://trends.google.com/trends/explore?hl=en&date=today%205-y&geo=BE&q=kasjmier%20sjaal%20%2B%20%C3%A9charpe%20cachemire%2Cwollen%20sjaal%20%2B%20%C3%A9charpe%20en%20laine%2Cwollen%20sokken%20%2B%20chaussettes%20laine%2Cwollen%20handschoenen%20%2B%20gants%20en%20laine%2Cwollen%20plaid%20%2B%20couverture%20en%20laine"),
    ("S38", "Google Trends DK Web Search five year core comparison", "https://trends.google.com/trends/explore?hl=en&date=today%205-y&geo=DK&q=cashmere%20halst%C3%B8rkl%C3%A6de%20%2B%20cashmere%20t%C3%B8rkl%C3%A6de%2Culd%20halst%C3%B8rkl%C3%A6de%20%2B%20uldt%C3%B8rkl%C3%A6de%2Culdsokker%2Culdhandsker%2Culdplaid%20%2B%20uldt%C3%A6ppe"),
    ("S39", "Google Trends SE Web Search five year core comparison", "https://trends.google.com/trends/explore?hl=en&date=today%205-y&geo=SE&q=kashmir%20halsduk%20%2B%20halsduk%20kashmir%2Cull%20halsduk%20%2B%20halsduk%20ull%2Cullsockor%20%2B%20ullstrumpor%2Cullhandskar%2Cullpl%C3%A4d%20%2B%20ullfilt"),
    ("S40", "Google Trends FI Web Search five year core comparison", "https://trends.google.com/trends/explore?hl=en&date=today%205-y&geo=FI&q=kashmirhuivi%20%2B%20cashmere%20huivi%2Cvillahuivi%2Cvillasukat%2Cvillak%C3%A4sineet%2Cvillahuopa"),
    ("S41", "Google Trends PL Web Search five year core comparison", "https://trends.google.com/trends/explore?hl=en&date=today%205-y&geo=PL&q=szalik%20kaszmirowy%2Cszalik%20we%C5%82niany%2Cskarpety%20we%C5%82niane%2Cr%C4%99kawiczki%20we%C5%82niane%2Ckoc%20we%C5%82niany"),
    ("S42", "Google Trends DE Google Shopping five year Wolldecke", "https://trends.google.com/trends/explore?hl=en&date=today%205-y&geo=DE&q=Wolldecke&gprop=froogle"),
    ("S43", "Google Trends DE YouTube Search five year Wolldecke", "https://trends.google.com/trends/explore?hl=en&date=today%205-y&geo=DE&q=Wolldecke&gprop=youtube"),
    ("S44", "Google Trends Help related and rising searches", "https://support.google.com/trends/answer/4355000?hl=en"),
    ("S45", "Google Trends Help interest by region", "https://support.google.com/trends/answer/4355212?hl=en"),
]

for sid, title, url in sources:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.03
    r = p.add_run(f"{sid}  ")
    set_run_font(r, size=8.5, bold=True, color=NAVY)
    add_hyperlink(p, title, url)

add_heading(doc, "Ограничения", 2)
add_bullets(doc, [
    "Google Trends — sampled normalized series. Ноль означает низкое разрешение в данном panel, а не доказанный нулевой спрос.",
    "TikTok search и public Instagram — snapshots, а не репрезентативные audience panels. Visible likes, comments и followers не раскрывают reach или sales.",
    "Meta Ad Library показывает delivery, targeting и creative, но не коммерческий spend, CAC, CVR или ROAS для этих commercial ads.",
    "HYPD absolute keyword estimates и Trends измеряют разные сущности; расхождения сохранены как data-quality warning, а не усреднены.",
    "Все финансовые выводы основаны на текущих landed cost и shipping assumptions из master. Фактические contracted tariffs, return incidence и payment mix могут изменить решение.",
])

add_heading(doc, "Итог", 2)
add_para(doc, "Следующий рациональный шаг — не ещё один broad competitor scan. До 21 сентября нужно провести Gate 0, выбрать физически доказанный hero и только затем собрать DE/NL cohorts. Текущая рабочая архитектура: cashmere scarf как conditional acquisition hero, mixed gift bundle как AOV layer, selected throws как organic and limited-drop lane, socks как pack or attach without cold paid. Любой другой порядок должен быть обоснован изменением freight, цены, качества или доступного stock.")

# Core properties
props = doc.core_properties
props.title = "Avdar Market исследование спроса и сценариев запуска"
props.subject = "Google Trends TikTok Instagram и экономика ассортимента"
props.author = "OpenAI Codex для Avdar Market"
props.keywords = "Avdar, Google Trends, TikTok, Instagram, cashmere, wool, ecommerce"
props.comments = "Срез данных 14 сентября 2026 года"

doc.save(OUT)
print(OUT)
